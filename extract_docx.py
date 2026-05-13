import docx
import sys

def read_docx(path):
    doc = docx.Document(path)
    text = ""
    for p in doc.paragraphs:
        if p.text.strip():
            text += p.text + "\n"
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text.strip() + " | "
            text += "\n"
        text += "-"*20 + "\n"
    return text

print(read_docx("Documentos/Casos de Uso INCASOFT Solutions..docx")[:5000])
