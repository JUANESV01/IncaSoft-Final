import json
import mimetypes
import os
from django.utils import timezone

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.core.exceptions import ImproperlyConfigured, ValidationError
from django.db.models import Count, Q, Sum
from django.db.models.functions import TruncMonth
from django.http import FileResponse, Http404, HttpRequest, HttpResponse, JsonResponse, QueryDict
from django.shortcuts import get_object_or_404, redirect, render
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from django.urls import reverse
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_POST

from .forms import (
    ColaboradorForm,
    DocumentoForm,
    EstadoIncapacidadForm,
    IncapacidadFiltroForm,
    IncapacidadForm,
    UsuarioActualizarForm,
    UsuarioCrearForm,
)
from .gemini_service import construir_resumen_agregado, generar_analisis_gemini, analizar_documento_incapacidad
from .models import Auditoria, Colaborador, Documento, HistorialEstado, Incapacidad, TipoIncapacidad
from .permissions import puede_editar_operacion, puede_gestionar_usuarios


def registrar_auditoria(usuario, accion, detalle=""):
    Auditoria.objects.create(usuario=usuario if usuario.is_authenticated else None, accion=accion, detalle=detalle)


def _incapacidades_filtradas(request):
    form = IncapacidadFiltroForm(request.GET or None)
    incapacidades = (
        Incapacidad.objects.select_related("colaborador", "tipo")
        .annotate(n_documentos=Count("documentos", distinct=True))
        .all()
    )
    if form.is_valid():
        q = form.cleaned_data.get("q")
        estado = form.cleaned_data.get("estado")
        tipo = form.cleaned_data.get("tipo")
        fecha_inicio = form.cleaned_data.get("fecha_inicio")
        fecha_fin = form.cleaned_data.get("fecha_fin")
        if q:
            incapacidades = incapacidades.filter(
                Q(colaborador__nombres__icontains=q)
                | Q(colaborador__apellidos__icontains=q)
                | Q(colaborador__numero_identificacion__icontains=q)
                | Q(numero_radicado__icontains=q)
            )
        if estado:
            incapacidades = incapacidades.filter(estado=estado)
        if tipo:
            incapacidades = incapacidades.filter(tipo=tipo)
        if fecha_inicio:
            incapacidades = incapacidades.filter(fecha_inicio__gte=fecha_inicio)
        if fecha_fin:
            incapacidades = incapacidades.filter(fecha_fin__lte=fecha_fin)
    return form, incapacidades


def _incapacidades_filtradas_desde_querystring(query_string: str):
    """Replica los filtros GET del reporte (misma lógica que la página de reportes)."""
    qs = (query_string or "").strip()
    if qs.startswith("?"):
        qs = qs[1:]
    shim = HttpRequest()
    shim.method = "GET"
    shim.GET = QueryDict(qs, mutable=True) if qs else QueryDict()
    return _incapacidades_filtradas(shim)


@login_required
def dashboard(request):
    from django.utils import timezone
    import json
    
    # Obtener todas las incapacidades con relaciones
    incapacidades = Incapacidad.objects.select_related("colaborador", "tipo").all()
    total = incapacidades.count()
    
    # Conteo por estado de forma segura
    conteo_estado = {estado: 0 for estado, _label in Incapacidad.ESTADO_CHOICES}
    if total > 0:
        for item in incapacidades.values("estado").annotate(total=Count("id")):
            conteo_estado[item["estado"]] = item["total"]

    # Métricas profesionales con fallback seguro
    hoy = timezone.now()
    inicio_mes = hoy.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    
    activas_count = incapacidades.exclude(estado__in=[Incapacidad.ESTADO_PAGADA, Incapacidad.ESTADO_RECHAZADA]).count()
    pendientes_count = incapacidades.filter(estado=Incapacidad.ESTADO_RECIBIDA).count()
    
    dias_query = incapacidades.filter(fecha_inicio__gte=inicio_mes.date()).aggregate(total=Sum("dias"))
    dias_mes = dias_query["total"] or 0
    
    tasa_tramite = 0
    if total > 0:
        tasa_tramite = int(((total - pendientes_count) / total) * 100)

    # Datos para gráficos (Blindaje total)
    labels_list = [str(label) for _value, label in Incapacidad.ESTADO_CHOICES]
    data_list = [conteo_estado.get(value, 0) for value, _label in Incapacidad.ESTADO_CHOICES]

    contexto = {
        "total_incapacidades": total,
        "total_colaboradores": Colaborador.objects.count(),
        "activas_count": activas_count,
        "pendientes_count": pendientes_count,
        "dias_mes": dias_mes,
        "tasa_tramite": tasa_tramite,
        "grafico_estados_labels": json.dumps(labels_list),
        "grafico_estados_data": json.dumps(data_list),
        "recientes": incapacidades.order_by("-fecha_creacion")[:6],
        "historial": HistorialEstado.objects.select_related("incapacidad", "usuario", "incapacidad__colaborador").order_by("-fecha")[:6],
        "tipos": TipoIncapacidad.objects.annotate(total=Count("incapacidades")).order_by("-total")[:5],
    }
    return render(request, "gestion/dashboard.html", contexto)


@login_required
def incapacidad_exportar_excel(request):
    """Exporta la lista filtrada de incapacidades a un archivo Excel (.xlsx)."""
    form, incapacidades = _incapacidades_filtradas(request)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incapacidades"
    
    # Encabezados
    headers = [
        "Radicado", "Colaborador", "Identificación", "Tipo", "Entidad", 
        "Inicio", "Fin", "Días Totales", "Días Empresa", "Días Entidad", "Estado"
    ]
    ws.append(headers)
    
    # Estilo encabezados
    header_fill = PatternFill(start_color="1A5F6E", end_color="1A5F6E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Datos
    for inc in incapacidades:
        ws.append([
            inc.numero_radicado,
            inc.colaborador.nombre_completo,
            inc.colaborador.numero_identificacion,
            inc.tipo.nombre,
            inc.get_entidad_responsable_display(),
            inc.fecha_inicio.strftime("%d/%m/%Y"),
            inc.fecha_fin.strftime("%d/%m/%Y"),
            inc.dias,
            inc.dias_empresa,
            inc.dias_entidad,
            inc.get_estado_display()
        ])

    # Ajuste de columnas
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column_letter].width = max_length + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="Reporte_Incapacidades.xlsx"'
    wb.save(response)
    return response


@login_required
def incapacidad_lista(request):
    form, incapacidades = _incapacidades_filtradas(request)
    return render(
        request,
        "gestion/incapacidad_lista.html",
        {
            "form": form,
            "incapacidades": incapacidades,
            "puede_editar": puede_editar_operacion(request.user),
        },
    )


@login_required
def incapacidad_crear(request):
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol tiene acceso de consulta, no de registro.")
        return redirect("gestion:incapacidad_lista")
    if request.method == "POST":
        form = IncapacidadForm(request.POST, request.FILES)
        if form.is_valid():
            incapacidad = form.save(commit=False)
            incapacidad.estado = Incapacidad.ESTADO_RECIBIDA
            incapacidad.creado_por = request.user
            incapacidad.actualizado_por = request.user
            try:
                incapacidad.full_clean()
                incapacidad.save()
                HistorialEstado.objects.create(
                    incapacidad=incapacidad,
                    estado_anterior=Incapacidad.ESTADO_RECIBIDA,
                    estado_nuevo=Incapacidad.ESTADO_RECIBIDA,
                    comentario="Registro inicial de la incapacidad.",
                    usuario=request.user,
                )
                registrar_auditoria(request.user, "Registro de incapacidad", str(incapacidad))
                messages.success(request, "Incapacidad registrada en estado Recibida.")
                return redirect(incapacidad)
            except ValidationError as exc:
                if hasattr(exc, "message_dict"):
                    for field, errors in exc.message_dict.items():
                        for error in errors:
                            form.add_error(field if field in form.fields else None, error)
                else:
                    form.add_error(None, exc)
    else:
        form = IncapacidadForm()
    return render(request, "gestion/incapacidad_formulario.html", {"form": form, "titulo": "Registrar incapacidad"})


@login_required
@require_POST
def incapacidad_analizar_documento(request):
    """Endpoint AJAX para analizar un documento con Gemini."""
    if not request.FILES.get("archivo"):
        return JsonResponse({"error": "No se subió ningún archivo"}, status=400)
    
    archivo = request.FILES["archivo"]
    content = archivo.read()
    mime = mimetypes.guess_type(archivo.name)[0] or "application/octet-stream"
    
    try:
        resultado = analizar_documento_incapacidad(
            content,
            mime,
            api_key=settings.GEMINI_API_KEY
        )
        return JsonResponse(resultado)
    except Exception as e:
        return JsonResponse({"error": str(e)}, status=500)


@login_required
def incapacidad_detalle(request, pk):
    incapacidad = get_object_or_404(
        Incapacidad.objects.select_related("colaborador", "tipo", "creado_por", "actualizado_por").prefetch_related(
            "documentos"
        ),
        pk=pk,
    )
    return render(
        request,
        "gestion/incapacidad_detalle.html",
        {
            "incapacidad": incapacidad,
            "historial": incapacidad.historial.select_related("usuario"),
            "puede_editar": puede_editar_operacion(request.user),
        },
    )


@login_required
def incapacidad_estado(request, pk):
    incapacidad = get_object_or_404(Incapacidad, pk=pk)
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol no permite actualizar estados.")
        return redirect(incapacidad)
    if request.method == "POST":
        form = EstadoIncapacidadForm(incapacidad, request.POST)
        if form.is_valid():
            try:
                incapacidad.cambiar_estado(
                    form.cleaned_data["estado"],
                    usuario=request.user,
                    comentario=form.cleaned_data["comentario"],
                )
                registrar_auditoria(request.user, "Actualizacion de estado", str(incapacidad))
                messages.success(request, "Estado actualizado y registrado en historial.")
                return redirect(incapacidad)
            except ValidationError as exc:
                form.add_error("estado", exc)
    else:
        form = EstadoIncapacidadForm(incapacidad)
    return render(
        request,
        "gestion/formulario.html",
        {"form": form, "titulo": f"Actualizar estado: {incapacidad.colaborador.nombre_completo}"},
    )


@login_required
def colaborador_lista(request):
    q = request.GET.get("q", "")
    doc = request.GET.get("doc", "")
    area = request.GET.get("area", "")
    colaboradores = Colaborador.objects.all()
    if q:
        colaboradores = colaboradores.filter(
            Q(nombres__icontains=q) | Q(apellidos__icontains=q)
        )
    if doc:
        colaboradores = colaboradores.filter(numero_identificacion__icontains=doc)
    if area:
        colaboradores = colaboradores.filter(area__icontains=area)
    return render(
        request,
        "gestion/colaborador_lista.html",
        {
            "colaboradores": colaboradores,
            "q": q,
            "doc": doc,
            "area": area,
            "puede_editar": puede_editar_operacion(request.user),
        },
    )


@login_required
def colaborador_crear(request):
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol tiene acceso de consulta, no de registro.")
        return redirect("gestion:colaborador_lista")
    if request.method == "POST":
        form = ColaboradorForm(request.POST)
        if form.is_valid():
            colaborador = form.save()
            registrar_auditoria(request.user, "Registro de colaborador", str(colaborador))
            messages.success(request, "Colaborador registrado.")
            return redirect(colaborador)
    else:
        form = ColaboradorForm()
    return render(request, "gestion/formulario.html", {"form": form, "titulo": "Registrar colaborador"})


@login_required
def colaborador_editar(request, pk):
    colaborador = get_object_or_404(Colaborador, pk=pk)
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol no permite actualizar colaboradores.")
        return redirect(colaborador)
    if request.method == "POST":
        form = ColaboradorForm(request.POST, instance=colaborador)
        if form.is_valid():
            colaborador = form.save()
            registrar_auditoria(request.user, "Actualizacion de colaborador", str(colaborador))
            messages.success(request, "Colaborador actualizado.")
            return redirect(colaborador)
    else:
        form = ColaboradorForm(instance=colaborador)
    return render(request, "gestion/formulario.html", {"form": form, "titulo": "Actualizar colaborador"})


@login_required
def colaborador_detalle(request, pk):
    colaborador = get_object_or_404(Colaborador, pk=pk)
    incapacidades = colaborador.incapacidades.select_related("tipo").all()
    return render(
        request,
        "gestion/colaborador_detalle.html",
        {
            "colaborador": colaborador,
            "incapacidades": incapacidades,
            "puede_editar": puede_editar_operacion(request.user),
        },
    )


@login_required
@ensure_csrf_cookie
def reportes(request):
    form, incapacidades = _incapacidades_filtradas(request)
    por_estado = list(incapacidades.values("estado").annotate(total=Count("id")).order_by("estado"))
    por_tipo = list(incapacidades.values("tipo__nombre").annotate(total=Count("id")).order_by("-total"))
    por_mes = [
        {
            "mes": item["mes_fecha"].strftime("%Y-%m") if item["mes_fecha"] else "Sin fecha",
            "total": item["total"],
        }
        for item in incapacidades.annotate(mes_fecha=TruncMonth("fecha_inicio"))
        .values("mes_fecha")
        .annotate(total=Count("id"))
        .order_by("mes_fecha")
    ]
    contexto = {
        "form": form,
        "incapacidades": incapacidades[:50],
        "total": incapacidades.count(),
        "dias": incapacidades.aggregate(total=Sum("dias"))["total"] or 0,
        "por_estado": por_estado,
        "por_tipo": por_tipo,
        "por_mes": por_mes,
        "gemini_ia_habilitada": bool(getattr(settings, "GEMINI_API_KEY", "")),
    }
    return render(request, "gestion/reportes.html", contexto)


@login_required
def reportes_pdf(request):
    """Vista de impresión/exportación PDF del reporte (abre en nueva pestaña)."""
    from django.utils import timezone as tz
    form, incapacidades = _incapacidades_filtradas(request)
    por_estado = list(incapacidades.values("estado").annotate(total=Count("id")).order_by("estado"))
    por_tipo = list(incapacidades.values("tipo__nombre").annotate(total=Count("id")).order_by("-total"))
    contexto = {
        "form": form,
        "incapacidades": incapacidades,
        "total": incapacidades.count(),
        "dias": incapacidades.aggregate(total=Sum("dias"))["total"] or 0,
        "por_estado": por_estado,
        "por_tipo": por_tipo,
        "ahora": tz.localtime(),
    }
    return render(request, "gestion/reportes_pdf.html", contexto)


@login_required
@require_POST
def reportes_ia_analisis(request):
    """Genera texto de análisis con Gemini a partir de agregados del reporte (sin PII)."""
    if not getattr(settings, "GEMINI_API_KEY", ""):
        return JsonResponse(
            {
                "ok": False,
                "error": "Configure GEMINI_API_KEY en el servidor (p. ej. en .env). Obtenga una clave en Google AI Studio.",
            },
            status=503,
        )

    try:
        payload = json.loads(request.body.decode() or "{}")
    except json.JSONDecodeError:
        return JsonResponse({"ok": False, "error": "Solicitud inválida."}, status=400)

    query_string = payload.get("query", "")
    if not isinstance(query_string, str):
        query_string = ""

    form, incapacidades = _incapacidades_filtradas_desde_querystring(query_string)
    if not form.is_valid():
        return JsonResponse({"ok": False, "error": "Los filtros enviados no son válidos."}, status=400)

    total = incapacidades.count()
    dias = incapacidades.aggregate(s=Sum("dias"))["s"] or 0
    por_estado = list(incapacidades.values("estado").annotate(total=Count("id")).order_by("estado"))
    por_tipo = list(incapacidades.values("tipo__nombre").annotate(total=Count("id")).order_by("-total"))
    por_mes = [
        {
            "mes": item["mes_fecha"].strftime("%Y-%m") if item["mes_fecha"] else "Sin fecha",
            "total": item["total"],
        }
        for item in incapacidades.annotate(mes_fecha=TruncMonth("fecha_inicio"))
        .values("mes_fecha")
        .annotate(total=Count("id"))
        .order_by("mes_fecha")
    ]
    estado_labels = dict(Incapacidad.ESTADO_CHOICES)
    resumen = construir_resumen_agregado(
        total=total,
        dias=dias,
        por_estado=por_estado,
        por_tipo=por_tipo,
        por_mes=por_mes,
        estado_labels=estado_labels,
    )

    try:
        texto = generar_analisis_gemini(
            resumen,
            api_key=settings.GEMINI_API_KEY,
            model_name=settings.GEMINI_MODEL,
        )
    except (ValueError, ImproperlyConfigured) as exc:
        return JsonResponse({"ok": False, "error": str(exc)}, status=502)

    registrar_auditoria(
        request.user,
        "Analisis IA (Gemini) en reportes",
        f"total={total}; modelo={settings.GEMINI_MODEL}",
    )
    return JsonResponse({"ok": True, "texto": texto})


@login_required
@user_passes_test(puede_gestionar_usuarios)
def usuario_lista(request):
    usuarios = User.objects.prefetch_related("groups").order_by("username")
    return render(request, "gestion/usuario_lista.html", {"usuarios": usuarios})


@login_required
@user_passes_test(puede_gestionar_usuarios)
def usuario_crear(request):
    if request.method == "POST":
        form = UsuarioCrearForm(request.POST)
        if form.is_valid():
            usuario = form.save()
            registrar_auditoria(request.user, "Creacion de usuario", usuario.username)
            messages.success(request, "Usuario creado. El rol quedará activo en el próximo inicio de sesión.")
            return redirect(reverse("gestion:usuario_lista"))
    else:
        form = UsuarioCrearForm()
    return render(request, "gestion/formulario.html", {"form": form, "titulo": "Crear usuario"})


@login_required
@user_passes_test(puede_gestionar_usuarios)
def usuario_editar(request, pk):
    usuario = get_object_or_404(User, pk=pk)

    # Protección absoluta para el admin
    if usuario.username == 'admin':
        messages.error(request, "El usuario administrador principal no puede ser editado por seguridad del sistema.")
        return redirect('gestion:usuario_lista')

    if request.method == "POST":
        form = UsuarioActualizarForm(request.POST, instance=usuario)
        if form.is_valid():
            form.save()
            registrar_auditoria(request.user, "Actualizacion de usuario", usuario.username)
            messages.success(request, "Usuario actualizado. El rol quedará activo en el próximo inicio de sesión.")
            return redirect(reverse("gestion:usuario_lista"))
    else:
        form = UsuarioActualizarForm(instance=usuario)
    return render(request, "gestion/formulario.html", {"form": form, "titulo": f"Editar usuario: {usuario.username}"})


# ── Documentos adjuntos ───────────────────────────────────────────────────────

@login_required
def documento_subir(request, incapacidad_pk):
    """Sube un archivo (PDF, Word o imagen) adjunto a una incapacidad."""
    incapacidad = get_object_or_404(Incapacidad, pk=incapacidad_pk)
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol no permite subir documentos.")
        return redirect(incapacidad)
    if request.method == "POST":
        form = DocumentoForm(request.POST, request.FILES)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.incapacidad = incapacidad
            doc.subido_por = request.user
            doc.save()
            registrar_auditoria(
                request.user,
                "Documento adjuntado",
                f"{doc.get_tipo_display()} - {incapacidad}",
            )
            messages.success(request, "Documento adjuntado correctamente.")
            return redirect(incapacidad)
    else:
        form = DocumentoForm()
    return render(
        request,
        "gestion/documento_subir.html",
        {
            "form": form,
            "incapacidad": incapacidad,
            "titulo": "Adjuntar documento",
        },
    )


@login_required
def documento_eliminar(request, pk):
    """Elimina un documento adjunto (solo POST)."""
    doc = get_object_or_404(Documento, pk=pk)
    incapacidad = doc.incapacidad
    if not puede_editar_operacion(request.user):
        messages.error(request, "Tu rol no permite eliminar documentos.")
        return redirect(incapacidad)
    if request.method == "POST":
        nombre = str(doc)
        doc.archivo.delete(save=False)   # borra el archivo físico
        doc.delete()
        registrar_auditoria(request.user, "Documento eliminado", nombre)
        messages.success(request, "Documento eliminado.")
    return redirect(incapacidad)


# ── Archivos: servir con sesión iniciada (no depender solo de /media/ en producción) ──


@login_required
def incapacidad_soporte_ver(request, pk):
    incapacidad = get_object_or_404(Incapacidad, pk=pk)
    if not incapacidad.soporte_medico:
        raise Http404("No hay soporte médico adjunto.")
    f = incapacidad.soporte_medico
    nombre = os.path.basename(f.name)
    try:
        fh = f.open("rb")
    except FileNotFoundError as exc:
        raise Http404("El archivo ya no está disponible en el servidor.") from exc
    content_type, _ = mimetypes.guess_type(nombre)
    resp = FileResponse(fh, as_attachment=False, filename=nombre)
    if content_type:
        resp["Content-Type"] = content_type
    resp["Content-Disposition"] = f'inline; filename="{nombre}"'
    return resp


@login_required
def documento_archivo_ver(request, pk):
    doc = get_object_or_404(Documento, pk=pk)
    f = doc.archivo
    nombre = os.path.basename(f.name)
    try:
        fh = f.open("rb")
    except FileNotFoundError as exc:
        raise Http404("El archivo ya no está disponible en el servidor.") from exc
    content_type, _ = mimetypes.guess_type(nombre)
    resp = FileResponse(fh, as_attachment=False, filename=nombre)
    if content_type:
        resp["Content-Type"] = content_type
    resp["Content-Disposition"] = f'inline; filename="{nombre}"'
    return resp


@login_required
def incapacidad_soporte_descargar(request, pk):
    """Descarga el soporte médico como adjunto (Content-Disposition: attachment)."""
    incapacidad = get_object_or_404(Incapacidad, pk=pk)
    if not incapacidad.soporte_medico:
        raise Http404("No hay soporte médico adjunto.")
    f = incapacidad.soporte_medico
    nombre = os.path.basename(f.name)
    try:
        fh = f.open("rb")
    except FileNotFoundError as exc:
        raise Http404("El archivo ya no está disponible en el servidor.") from exc
    content_type, _ = mimetypes.guess_type(nombre)
    resp = FileResponse(fh, as_attachment=True, filename=nombre)
    if content_type:
        resp["Content-Type"] = content_type
    return resp


@login_required
def documento_archivo_descargar(request, pk):
    doc = get_object_or_404(Documento, pk=pk)
    f = doc.archivo
    nombre = os.path.basename(f.name)
    try:
        fh = f.open("rb")
    except FileNotFoundError as exc:
        raise Http404("El archivo ya no está disponible en el servidor.") from exc
    content_type, _ = mimetypes.guess_type(nombre)
    resp = FileResponse(fh, as_attachment=True, filename=nombre)
    if content_type:
        resp["Content-Type"] = content_type
    return resp
